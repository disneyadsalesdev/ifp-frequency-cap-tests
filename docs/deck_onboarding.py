"""Onboarding slides — learn first, then install. Clear TOC and new-hire path."""

from __future__ import annotations

from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from pptx.presentation import Presentation


def section_divider(prs, title: str, subtitle: str) -> None:
    from build_deck import add_title_bar, blank_slide

    slide = blank_slide(prs)
    add_title_bar(slide, title, prs, subtitle)


def build_navigation(prs) -> None:
    from build_deck import add_bullets, add_table, add_title_bar, blank_slide, title_slide

    title_slide(
        prs,
        "Cursor 101",
        "New hire training — follow the Table of Contents in order",
        "Next: Table of Contents",
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Table of Contents", prs, "Use Ctrl+F in Slide Sorter to jump by section title")
    add_table(
        slide,
        ("Section", "Slides to find", "What you do"),
        [
            ("1", "Learn Cursor First", "What it is, UI, screenshots — before install"),
            ("2", "Download & Install Cursor", "cursor.com → install → sign in"),
            ("3", "Day 1 Setup — Folders & MCP", "Create folders, paste files, mcp.json"),
            ("3B", "Understand MCP (IFP)", "How Cursor talks to server.py → IFP API"),
            ("4", "Cursor Settings", "Open Settings, MCP, Run Mode"),
            ("5", "Chat, New Agent & Run Buttons", "Ctrl+I, New Chat, Agent, Run/Skip"),
            ("6", "New Hire Startup", "First prompts + day-to-day tool map"),
            ("7", "IFP Every Day", "VPN, forecast in chat"),
            ("8", "IFP DevTools Targeting Lab", "F12 → copy JSON from portal"),
            ("9", "MSTR Daily Report", "Only if this is your job"),
            ("10", "Mission Control", "Schedule MSTR emails in browser"),
        ],
        top=0.95,
        row_h=0.34,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "WHERE DO I START?", prs, "Pick your situation — go to that section")
    add_table(
        slide,
        ("I need to…", "Go to section"),
        [
            ("Understand what Cursor is", "Section 1 — Learn Cursor First"),
            ("Download and install Cursor", "Section 2 — Download & Install"),
            ("Fix MCP or folder setup", "Section 3 — Day 1 Setup (create folders & paste files)"),
            ("Understand how MCP works", "Section 3B — Understand MCP (IFP)"),
            ("Fix Settings or MCP not green", "Section 4 — Cursor Settings"),
            ("Learn New Chat, Agent, Run buttons", "Section 5 — Chat & Run Buttons"),
            ("First-day copy-paste prompts", "Section 6 — New Hire Startup"),
            ("Run a forecast today", "Section 7 — IFP Every Day"),
            ("Copy targeting from IFP website", "Section 8 — DevTools Lab"),
            ("Run the daily MSTR email script", "Section 9 — MSTR Daily Report"),
            ("Schedule or fix MSTR emails", "Section 10 — Mission Control"),
        ],
        top=1.0,
        row_h=0.32,
    )
    add_bullets(
        slide,
        ["", "Also: docs/NEW-USER-SIMPLE-GUIDE.pdf  ·  Appendix at end of deck for glossary"],
        top=4.55,
        size=12,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Day 1 Checklist", prs, "Check off before live IFP forecasts")
    add_bullets(
        slide,
        [
            "□ Section 1 — Read what Cursor is (with pictures)",
            "□ Section 2 — Download & install Cursor from cursor.com",
            "□ Section 3 — Create folders, paste JSON + server.py, MCP green",
            "□ Section 3B — Skim how MCP + stdio works (optional but helpful)",
            "□ Section 4 — Settings: Run Mode asks before commands",
            "□ Section 5 — Practice New Chat + Agent + Run button",
            "□ Section 6 — Run startup prompts (Ask then Agent)",
            "□ Section 7 — One test forecast (VPN on)",
        ],
        size=15,
    )


def build_learn_cursor(prs) -> None:
    from build_deck import (
        add_bullets,
        add_cursor_ui_mockup,
        add_gfg_image_slide,
        add_image_or_placeholder,
        add_table,
        add_title_bar,
        blank_slide,
    )

    section_divider(prs, "SECTION 1 — Learn Cursor First", "Read this BEFORE you install")

    slide = blank_slide(prs)
    add_title_bar(slide, "What Is Cursor?", prs)
    add_bullets(
        slide,
        [
            "An app on your computer — a smart notebook for your work files.",
            "Looks like VS Code. Chat on the side (Ctrl+I) — type plain English.",
            "Reads your folder and runs team tools when you approve.",
            "",
            "You do NOT need to be a programmer. Cursor asks before changing anything.",
        ],
        size=16,
    )

    add_gfg_image_slide(
        prs,
        "What Is Cursor AI?",
        "gfg-what-is-cursor.png",
        "AI-powered editor · built on VS Code",
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Key Features & Modes", prs)
    add_table(
        slide,
        ("Feature / mode", "Plain English"),
        [
            ("Chat (Ctrl+I)", "Talk to the helper"),
            ("New Chat (+)", "Fresh conversation — use for every new task"),
            ("Ask mode", "Explains only — safe to practice"),
            ("Agent mode", "Runs IFP, edits files — with your OK"),
            ("@ filename", "Point at one file in chat"),
            ("MCP tools", "Extra buttons like Run IFP forecast"),
        ],
        top=1.1,
        row_h=0.4,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "The Cursor Window", prs)
    add_cursor_ui_mockup(slide, top=1.0)
    add_table(
        slide,
        ("Part", "What it is"),
        [
            ("Left", "File list — your project folder"),
            ("Center", "File you are viewing"),
            ("Right", "Chat panel (Ctrl+I)"),
            ("Bottom", "Terminal — commands run here with your approval"),
        ],
        top=4.85,
        row_h=0.32,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Pictures to Recognize Later", prs, "Capture after Section 6 — save in docs/images/")
    add_image_or_placeholder(slide, "00-ifp-open-folder.png", 0.55, 1.0, 3.0, 1.5, "Open Folder")
    add_image_or_placeholder(slide, "02-agent-chat.png", 3.7, 1.0, 3.0, 1.5, "Agent mode")
    add_image_or_placeholder(slide, "00-ifp-diff-accept.png", 6.85, 1.0, 2.5, 1.5, "Accept/Reject")
    add_image_or_placeholder(slide, "00-ifp-ask-cap-json.png", 0.55, 2.7, 4.0, 1.5, "Ask chat")
    add_image_or_placeholder(slide, "00-ifp-forecast.png", 4.7, 2.7, 4.65, 1.5, "Forecast (optional)")


def build_part_a_install(prs) -> None:
    from build_deck import add_bullets, add_title_bar, blank_slide

    section_divider(prs, "SECTION 2 — Download & Install Cursor", "Part A — every click")

    slide = blank_slide(prs)
    add_title_bar(slide, "How to Download Cursor", prs, "Do this on your work PC")
    add_bullets(
        slide,
        [
            "1. Open Chrome or Edge",
            "2. Click the address bar at the top → type: cursor.com → press Enter",
            "3. On the Cursor home page, click the big Download button",
            "4. Choose Windows (or Mac if you are on a Mac)",
            "5. Wait for the download to finish",
            "     Usually saves as: Downloads\\Cursor Setup.exe",
            "",
            "Tip: If IT blocks the site, ask your trainer for the installer.",
        ],
        size=15,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "How to Install & Open Cursor", prs)
    add_bullets(
        slide,
        [
            "6. Open File Explorer → go to Downloads",
            "7. Double-click Cursor Setup.exe",
            "8. If Windows asks “Allow this app?” → click Yes",
            "9. In the installer: Next → Install → Finish",
            "10. Start menu → type Cursor → open the app",
            "11. Sign in if prompted (use work account if IT requires SSO)",
            "12. “Import VS Code settings?” → Skip or Continue (either is fine)",
            "13. Optional: Help → Check for Updates · pin Cursor to taskbar",
        ],
        size=14,
    )


def build_part_b_folders_mcp(prs) -> None:
    from build_deck import add_bullets, add_code_block, add_table, add_title_bar, blank_slide

    section_divider(
        prs,
        "SECTION 3 — Day 1 Setup: Folders & MCP",
        "Part B — create folders yourself and paste team files",
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Folder Layout on Your PC", prs, "Create this tree under C:\\Users\\<you>\\projects\\")
    add_bullets(
        slide,
        [
            "projects\\",
            "  ifp-frequency-cap-tests\\     ← open this daily in Cursor",
            "    config\\base-request.json",
            "    reference\\dma-codes.json",
            "    reference\\cap-ratio-expectations.json",
            "    scripts\\   (optional — can stay empty at first)",
            "    output\\    (empty — results go here later)",
            "  ifp-mcp-server\\",
            "    server.py",
            "    requirements.txt",
            "",
            "Also: C:\\Users\\<you>\\.cursor\\mcp.json",
        ],
        size=13,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Step 1 — Create the IFP Project Folders", prs)
    add_bullets(
        slide,
        [
            "1. File Explorer → C:\\Users\\<your name>\\ → create folder: projects",
            "2. Inside projects → create folder: ifp-frequency-cap-tests",
            "3. Inside ifp-frequency-cap-tests, create these subfolders:",
            "     config",
            "     reference",
            "     scripts  (optional for day 1)",
            "     output   (leave empty)",
            "4. Sibling folder: projects\\ifp-mcp-server  (same level as step 2)",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Step 2 — config/base-request.json", prs, "The forecast request template the API sends")
    add_bullets(
        slide,
        [
            "1. Inside config\\ → create file: base-request.json",
            "2. Paste JSON that includes:",
            "     start-date, end-date, ad-products",
            "     frequency-cap-detail (cap limit, tier, product-id)",
            "     targeting-detail (publisher, country, dma-code, etc.)",
            "",
            "Where to get the first paste:",
            "  • Best: Section 8 DevTools — copy Request JSON from IFP portal",
            "  • Or: copy from a teammate’s config\\base-request.json",
            "  • Or: team shared drive / wiki starter template",
            "",
            "Save as UTF-8 text. File must be valid JSON (matching braces).",
        ],
        size=12,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Step 3 — reference/ JSON Files", prs, "Lookup and test files — paste team copies")
    add_table(
        slide,
        ("Create this file", "What to paste", "Used for"),
        [
            (
                "reference\\dma-codes.json",
                "Team lookup: city name ↔ DMA number (by_name, by_code)",
                "Ask chat: “DMA code for Baton Rouge?”",
            ),
            (
                "reference\\cap-ratio-expectations.json",
                "Team list of cap test cases + expected ratios",
                "Cap test matrix / validation",
            ),
        ],
        top=1.05,
        row_h=0.55,
    )
    add_bullets(
        slide,
        [
            "",
            "Copy both files from a teammate or team shared location.",
            "Do not send these whole files to the API — MCP reads them as lookups.",
        ],
        top=3.35,
        size=13,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Step 4 — ifp-mcp-server Files", prs, "Two files Cursor needs to start MCP")
    add_bullets(
        slide,
        [
            "Inside projects\\ifp-mcp-server\\ create:",
            "",
            "A) requirements.txt — paste exactly this one line:",
        ],
        size=14,
    )
    add_code_block(slide, "mcp>=1.9.0", 2.0)
    add_bullets(
        slide,
        [
            "",
            "B) server.py — copy from a teammate OR build using Section 3B",
            "     (Official guide: modelcontextprotocol.io → Develop → Build Server)",
            "     You must have server.py in this folder — not a shortcut.",
            "",
            "Do not edit server.py on day 1 unless your trainer asks.",
        ],
        top=2.45,
        size=12,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Step 5 — Python & pip", prs)
    add_bullets(
        slide,
        [
            "1. Install Python 3 (python.org or IT) → PowerShell: py --version",
            "2. py -m pip install -r C:\\Users\\<you>\\projects\\ifp-mcp-server\\requirements.txt",
            "     (change <you> to your Windows username)",
            "3. C:\\Users\\<you>\\ → View → Hidden items → open or create .cursor\\",
            "4. Next slide — create mcp.json in that folder",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(
        slide,
        "Step 5 — mcp.json (Copy & Paste)",
        prs,
        "Select all text below → paste into C:\\Users\\<you>\\.cursor\\mcp.json",
    )
    from build_deck import add_code_block

    mcp_json = (
        '{\n'
        '  "mcpServers": {\n'
        '    "ifp-forecast": {\n'
        '      "type": "stdio",\n'
        '      "command": "py",\n'
        '      "args": ["C:/Users/YourName/projects/ifp-mcp-server/server.py"],\n'
        '      "env": {\n'
        '        "IFP_API_URL": "http://inventory-forecasting-prod-dplus.ava.prod.hulu.com/v1/inventory/forecast",\n'
        '        "IFP_SOURCE_HEADER": "RYM Frequency Cap Test",\n'
        '        "IFP_TESTS_ROOT": "C:/Users/YourName/projects/ifp-frequency-cap-tests"\n'
        '      }\n'
        '    }\n'
        '  }\n'
        '}'
    )
    add_bullets(
        slide,
        [
            "Replace every YourName with your Windows username before saving.",
            "Or copy from: docs/templates/mcp.json in the IFP project folder.",
        ],
        top=1.0,
        size=12,
    )
    add_code_block(slide, mcp_json, 1.55, font_size=9, max_height=5.5)

    slide = blank_slide(prs)
    add_title_bar(slide, "Step 6 — Open Project & Verify MCP", prs)
    add_bullets(
        slide,
        [
            "1. Cursor → File → Open Folder → ifp-frequency-cap-tests",
            "2. Confirm left sidebar shows config\\ and reference\\ with your JSON files",
            "3. Ctrl+Shift+P → Reload Window (after saving mcp.json)",
            "4. Section 4 → Settings → MCP → ifp-forecast must be GREEN",
            "",
            "Red MCP? Missing server.py, wrong path, pip failed, or bad JSON path.",
            "More detail: docs/complete-ifp-mcp-setup-guide.md · Section 3B below",
        ],
        size=13,
    )


def build_mcp_explainer(prs) -> None:
    """Section 3B — MCP concepts from official docs, mapped to ifp-mcp-server."""
    from build_deck import add_bullets, add_code_block, add_table, add_title_bar, blank_slide

    section_divider(
        prs,
        "SECTION 3B — Understand MCP (IFP)",
        "From modelcontextprotocol.io — how our forecast server works",
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "MCP — Three Things a Server Can Offer", prs, "Official MCP docs — we mainly use Tools")
    add_table(
        slide,
        ("MCP capability", "Plain English", "IFP example"),
        [
            ("Resources", "Read-only data (like opening a file)", "Not used in our server"),
            ("Tools", "Actions Cursor runs — you approve first", "run_forecast, cap test matrix"),
            ("Prompts", "Pre-written task templates", "Not used in our server"),
        ],
        top=1.05,
        row_h=0.48,
    )
    add_bullets(
        slide,
        [
            "",
            "Cursor = MCP host (client)  ·  ifp-mcp-server/server.py = MCP server",
            "mcp.json tells Cursor how to start the server on your PC.",
        ],
        top=3.25,
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "What Happens When You Ask for a Forecast", prs)
    add_bullets(
        slide,
        [
            "1. You type in Cursor chat (Agent mode)",
            "2. Cursor picks the MCP tool run_forecast (if that fits your question)",
            "3. You click Run / Approve on the MCP approval card",
            "4. Cursor starts server.py using the command in mcp.json (stdio transport)",
            "5. server.py reads config/base-request.json from IFP_TESTS_ROOT",
            "6. server.py POSTs JSON to the IFP forecast API (VPN must be on)",
            "7. capacity / available / ratio come back to chat",
            "",
            "You did not run PowerShell commands yourself — the server did the API call.",
        ],
        size=13,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Our IFP MCP Tools", prs, "Defined in server.py with @mcp.tool()")
    add_table(
        slide,
        ("Tool name", "What it does"),
        [
            ("run_forecast", "One forecast — cap, dates, optional DMA codes"),
            ("run_cap_test_matrix", "Run all cases in cap-ratio-expectations.json"),
            ("validate_cap_test_matrix", "Same runs + PASS/FAIL vs expected ratios"),
            ("ifp_server_info", "Show paths, API URL, and tool list (good first test)"),
        ],
        top=1.1,
        row_h=0.42,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "mcp.json & stdio — Rules That Prevent Breakages", prs)
    add_bullets(
        slide,
        [
            "Transport: type stdio — Cursor talks to server.py over stdin/stdout",
            "",
            "Use ABSOLUTE paths in mcp.json (forward slashes OK on Windows)",
            "After any mcp.json edit: Ctrl+Shift+P → Reload Window",
            "",
            "Test manually (optional):",
            "  py C:\\...\\ifp-mcp-server\\server.py",
            "  If it looks frozen — that is NORMAL (waiting for Cursor). Ctrl+C to stop.",
            "",
            "If MCP is RED: Settings → MCP → ifp-forecast → read error log",
            "",
            "Builders: NEVER use print() in server.py — it breaks stdio JSON messages.",
            "  Use Python logging module (writes to stderr) instead.",
        ],
        size=12,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Building server.py — Official Pattern", prs, "Same structure as MCP weather tutorial — mapped to IFP")
    add_bullets(
        slide,
        [
            "Official tutorial: modelcontextprotocol.io/docs/develop/build-server",
            "",
            "Minimum pieces (Python + FastMCP — what our server uses):",
        ],
        size=13,
    )
    add_code_block(
        slide,
        "from mcp.server.fastmcp import FastMCP\n"
        "mcp = FastMCP(\"ifp-forecast\")\n\n"
        "@mcp.tool()\n"
        "def run_forecast(limit: int, duration: int, duration_unit: str) -> str:\n"
        '    """Run one IFP forecast — docstring shows in Cursor tool list."""\n'
        "    ... read base-request.json, call IFP API ...\n\n"
        "if __name__ == \"__main__\":\n"
        "    mcp.run()   # stdio — Cursor starts this, not you in daily use",
        2.0,
    )
    add_bullets(
        slide,
        [
            "",
            "Build checklist: folder → requirements.txt (mcp>=1.9.0) → server.py →",
            "pip install → mcp.json entry → Reload Window → MCP green → ifp_server_info in chat",
        ],
        top=5.35,
        size=12,
    )


def build_part_c_settings(prs) -> None:
    from build_deck import add_bullets, add_table, add_title_bar, blank_slide

    section_divider(prs, "SECTION 4 — Cursor Settings", "Part C — open Settings and verify MCP")

    slide = blank_slide(prs)
    add_title_bar(slide, "How to Open Cursor Settings", prs)
    add_bullets(
        slide,
        [
            "Method 1: Click the gear icon at the bottom-left of Cursor",
            "Method 2: Press Ctrl+, (comma) on your keyboard",
            "",
            "The Settings panel opens on the right or in a tab.",
            "Use the search box at the TOP of Settings to find options quickly.",
            "",
            "Settings you MUST check on day 1:",
            "  • Search: MCP  →  ifp-forecast status",
            "  • Search: Run Mode  (or Agents → Approvals)",
        ],
        size=15,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "MCP Settings — Must Be Green", prs)
    add_bullets(
        slide,
        [
            "1. Settings → search MCP",
            "2. Find ifp-forecast in the list",
            "3. Status should show GREEN / connected",
            "",
            "If RED:",
            "  • Click the server name → read the error log",
            "  • Fix paths in C:\\Users\\<you>\\.cursor\\mcp.json",
            "  • Ctrl+Shift+P → Reload Window",
            "",
            "mcp.json lives on YOUR PC only — never put passwords in it.",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Run Mode — Command Approvals", prs, "Settings → search Run Mode or Agents → Approvals")
    add_table(
        slide,
        ("Setting name", "What it means for new hires"),
        [
            ("Allowlist (empty)", "Safest — Cursor asks before almost every command/MCP run"),
            ("Auto-review", "Some runs happen after review — fewer popups"),
            ("Run everything", "NO prompts — do NOT use on day 1"),
        ],
        top=1.1,
        row_h=0.48,
    )
    add_bullets(
        slide,
        [
            "",
            "Recommendation: start with Allowlist (empty) or the mode that always asks you.",
            "After any mcp.json edit: Ctrl+Shift+P → Reload Window.",
        ],
        top=3.15,
        size=14,
    )


def build_part_d_chat(prs) -> None:
    from build_deck import add_bullets, add_table, add_title_bar, blank_slide

    section_divider(prs, "SECTION 5 — Chat, New Agent & Run Buttons", "Part D")

    slide = blank_slide(prs)
    add_title_bar(slide, "Open Chat & Start a New Agent", prs, "Do this for EVERY new task")
    add_bullets(
        slide,
        [
            "1. Press Ctrl+I → chat panel opens (usually on the right)",
            "2. At the TOP of the chat panel, click + or New Chat",
            "     Why? Old messages confuse the Agent on a new task.",
            "3. Click the mode dropdown at the top of chat:",
            "     • Ask — learning only, no file changes",
            "     • Agent — forecasts, scripts, file edits (needs your OK)",
            "     • Plan — big jobs; makes a to-do list first",
            "4. Model dropdown — leave team default unless IT says otherwise",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "When to Use New Chat + Agent", prs)
    add_table(
        slide,
        ("Your task", "New Chat?", "Mode"),
        [
            ("Explain a JSON file", "Yes", "Ask"),
            ("New forecast question", "Yes", "Agent"),
            ("Run daily report script", "Yes", "Agent"),
            ("Follow-up on SAME forecast", "No — same chat OK", "Agent"),
            ("Totally different job", "Yes — always", "Agent or Ask"),
        ],
        top=1.05,
        row_h=0.42,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Run, Skip & Allowlist Buttons", prs, "When Agent wants to run a command or MCP tool")
    add_table(
        slide,
        ("Button you see", "What to do"),
        [
            ("Run / Approve", "OK this command or MCP tool ONE time"),
            ("Skip / Reject", "Do NOT run — ask trainer if unsure"),
            ("Add to allowlist", "Trust this command later — ask trainer first"),
            ("Always allow / Always run", "Auto-approve this tool — ask trainer before using"),
        ],
        top=1.05,
        row_h=0.4,
    )
    add_bullets(
        slide,
        [
            "",
            "IFP forecasts: VPN must be ON before you click Run.",
            "These buttons are for terminal commands and MCP — NOT the same as file edits.",
        ],
        top=3.35,
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Accept & Reject — File Changes", prs, "Separate from Run / Skip")
    add_bullets(
        slide,
        [
            "When Agent edits a file, you see a diff (red = removed, green = added).",
            "",
            "Accept — keep the change",
            "Reject — undo the change",
            "",
            "Always read the diff before Accept.",
            'Not sure? Reject and type: "Explain that change first"',
        ],
        size=16,
    )


def build_new_hire_startup(prs) -> None:
    """Section 6 — initial prompts and day-to-day tool map."""
    from build_deck import add_bullets, add_code_block, add_table, add_title_bar, blank_slide

    section_divider(prs, "SECTION 6 — New Hire Startup", "Settings done? Run these prompts in order")

    slide = blank_slide(prs)
    add_title_bar(slide, "New Hire — What to Do First in Cursor", prs, "After Sections 2–5 are complete")
    add_bullets(
        slide,
        [
            "1. File → Open Folder → ifp-frequency-cap-tests",
            "2. Confirm Section 4: MCP green, Run Mode asks before run",
            "3. Section 6 prompts below — Ask first (safe), then Agent",
            "4. VPN on → one test forecast (Section 7)",
            "5. Screenshot key steps (Section 1 picture slide)",
            "",
            "Pair with your trainer for Steps 3–4 on day 1.",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Startup Prompts — Ask Mode (Steps 1–3)", prs, "New Chat → Ask for EACH — NO VPN")
    add_bullets(slide, ["Copy one prompt, send, then New Chat for the next:"], top=1.0, size=13)
    add_code_block(
        slide,
        "Step 1 — Explain reference/cap-ratio-expectations.json in simple terms.\n"
        "What is one test case?\n\n"
        "Step 2 — What is the DMA code for Baton Rouge?\n"
        "Use reference/dma-codes.json.\n\n"
        "Step 3 — Explain config/base-request.json — dates and targeting.\n"
        "Do not change any files.",
        1.35,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Startup Prompts — Agent (Steps 4–5)", prs)
    add_bullets(slide, ["Step 4 — New Chat → Agent:"], top=1.0, size=13)
    add_code_block(
        slide,
        "Create practice-notes.md with 4 bullets about caps, DMA codes,\n"
        "and base-request.json. Review diff → Accept.",
        1.35,
    )
    add_bullets(slide, ["Step 5 — VPN ON → New Chat → Agent → click Run if MCP asks:"], top=2.55, size=12)
    add_code_block(
        slide,
        "Run a forecast for July 8, 2026, Hulu US, Baton Rouge, 2 per hour.\n"
        "Show capacity and available.",
        2.95,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Day-to-Day Tools — Quick Map", prs, "Which folder / tool for which job")
    add_table(
        slide,
        ("Job", "Where / how", "VPN?"),
        [
            ("IFP forecast", "ifp-frequency-cap-tests → Agent + MCP Run", "Yes"),
            ("Cap test matrix", "Same folder → Agent: run cap test matrix", "Yes"),
            ("Edit targeting JSON", "config/base-request.json (Section 8 DevTools)", "To test"),
            ("Daily MSTR email", "all-publishers-daily-report → dailypublisherscript.py", "No"),
            ("Fix MCP", "C:\\Users\\<you>\\.cursor\\mcp.json → Reload Window", "No"),
            ("Schedule MSTR emails", "Mission Control in browser — Section 10", "Per org"),
        ],
        top=1.0,
        row_h=0.38,
    )


def build_ifp_daily(prs) -> None:
    from build_deck import add_bullets, add_code_block, add_table, add_title_bar, blank_slide

    section_divider(prs, "SECTION 7 — IFP Every Day", "Part F — repeat for each forecast")

    slide = blank_slide(prs)
    add_title_bar(slide, "IFP Forecast — Step by Step", prs)
    add_table(
        slide,
        ("Step", "Action"),
        [
            ("1", "Turn on Hulu VPN"),
            ("2", "Cursor → File → Open Folder → ifp-frequency-cap-tests"),
            ("3", "Ctrl+I → click New Chat (+) → select Agent mode"),
            ("4", "Type forecast in plain English (date, Hulu US, city, cap unit)"),
            ("5", "When Run/Approve appears for MCP → click Run"),
            ("6", "Read capacity and available in the reply"),
        ],
        top=1.05,
        row_h=0.4,
    )
    add_bullets(
        slide,
        [
            "",
            "Cap units: MINUTE, HOUR, DAY only (e.g. 4 per 30 DAY).",
            "Cities use DMA numbers — see reference/dma-codes.json or ask in chat.",
        ],
        top=3.55,
        size=13,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "IFP — Example Prompt", prs, "New Chat → Agent → VPN on")
    add_code_block(
        slide,
        "Run a forecast for 7/10 Hulu US, Baton Rouge, 2 per hour.\n"
        "Show capacity, available, and ratio.",
        1.4,
    )


def build_devtools_lab(prs) -> None:
    from build_deck import add_bullets, add_title_bar, blank_slide

    section_divider(
        prs,
        "SECTION 8 — IFP DevTools Targeting Lab",
        "When portal numbers ≠ API — copy exact JSON from the UI",
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "DevTools Lab — Portal Setup (G1–G3)", prs)
    add_bullets(
        slide,
        [
            "Goal: copy the exact JSON the IFP portal sends → config/base-request.json",
            "",
            "G1. Turn on Hulu VPN",
            "G2. Open browser → IFP portal → log in (SSO)",
            "G3. Build the forecast in the UI exactly like the deal:",
            "     • Dates  • Product (note product-id)  • Targeting  • DMA  • Frequency cap",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "DevTools Lab — Open Network Tab (G4–G7)", prs)
    add_bullets(
        slide,
        [
            "G4. Press F12 (or Ctrl+Shift+I) → Developer Tools panel opens",
            "G5. Click the Network tab",
            "G6. Check Preserve log (keeps requests after you click around)",
            "G7. Set filter to Fetch/XHR (hides images and CSS noise)",
        ],
        size=16,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "DevTools Lab — Capture Request JSON (G8–G12)", prs)
    add_bullets(
        slide,
        [
            "G8. In the portal, click Run forecast / Submit",
            "G9. In the Network list, find the row with forecast in the name or URL",
            "G10. Click that row",
            "G11. Right panel → Payload or Request tab (NOT Response)",
            "G12. Copy the full JSON request body",
        ],
        size=15,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "DevTools Lab — Save & Test (G13–G16)", prs)
    add_bullets(
        slide,
        [
            "G13. In Cursor, open config/base-request.json",
            "G14. Replace contents with copied JSON (or ask Agent to merge carefully)",
            "G15. Save the file (Ctrl+S)",
            "G16. VPN on → Section 7 forecast → compare numbers to portal",
            "",
            "WARNING: product-id in the UI does NOT auto-fill API targeting.",
            "JSON must include full targeting-detail — not product-id alone.",
        ],
        size=14,
    )


def build_mstr_and_mission_control(prs) -> None:
    from build_deck import add_bullets, add_table, add_title_bar, blank_slide

    section_divider(prs, "SECTION 9 — MSTR Daily Report", "Part H — only if this is your job")

    slide = blank_slide(prs)
    add_title_bar(slide, "Daily MSTR Report in Cursor", prs)
    add_table(
        slide,
        ("Step", "Action"),
        [
            ("1", "Wait for MicroStrategy email in Outlook (.xlsx attachments)"),
            ("2", "Cursor → File → Open Folder → all-publishers-daily-report"),
            ("3", "Ctrl+I → New Chat → Agent mode"),
            ("4", "Type: Run dailypublisherscript.py"),
            ("5", "Click Run if terminal approval appears"),
            ("6", "Review Outlook draft email → YOU click Send (Cursor does not send)"),
        ],
        top=1.05,
        row_h=0.42,
    )
    add_bullets(
        slide,
        ["", "No email arrived? Check Section 10 Mission Control schedule."],
        top=3.65,
        size=14,
    )

    section_divider(prs, "SECTION 10 — Mission Control", "Part I — browser tool, NOT Cursor — setup once")

    slide = blank_slide(prs)
    add_title_bar(slide, "Mission Control — Find the Report (I1–I5)", prs)
    add_bullets(
        slide,
        [
            "Mission Control schedules report emails to Outlook (set up once, verify weekly).",
            "",
            "I1. Ask manager for MicroStrategy Library URL + Mission Control link",
            "I2. Connect VPN/network per your org rules",
            "I3. Open MicroStrategy Library in browser → SSO login",
            "I4. Search: All Publisher Daily Report (or your team report name)",
            "I5. Open the report → confirm it looks correct",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Mission Control — Create Subscription (I6–I11)", prs)
    add_bullets(
        slide,
        [
            "I6. Open Mission Control (Library menu or direct link from manager)",
            "I7. Create subscription / New schedule",
            "I8. Select the report (All Publisher Daily Report)",
            "I9. Recurrence: Daily (before Section 9 script) or Weekly",
            "I10. Format: Excel (.xlsx) — daily script expects Excel attachments",
            "I11. Delivery: team distribution email list → Save",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Mission Control — Test & Weekly Check", prs)
    add_bullets(
        slide,
        [
            "I12. Run a test send → confirm email arrives in Outlook with .xlsx",
            "",
            "Each business day: did the MSTR email arrive? → run Section 9 script",
            "Weekly: did Mission Control job succeed? (ask analytics if emails stop)",
            "After IFP deal/targeting change: Section 8 DevTools → Section 7 test forecast",
        ],
        size=14,
    )


def build_full_onboarding(prs) -> None:
    build_navigation(prs)
    build_learn_cursor(prs)
    build_part_a_install(prs)
    build_part_b_folders_mcp(prs)
    build_mcp_explainer(prs)
    build_part_c_settings(prs)
    build_part_d_chat(prs)
    build_new_hire_startup(prs)
    build_ifp_daily(prs)
    build_devtools_lab(prs)
    build_mstr_and_mission_control(prs)
