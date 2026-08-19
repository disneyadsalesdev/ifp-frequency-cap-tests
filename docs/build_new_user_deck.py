"""Build NEW-USER-SIMPLE-GUIDE.pptx — simple onboarding slides (Windows + Mac)."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

DOCS_DIR = Path(__file__).parent
OUTPUT = DOCS_DIR / "NEW-USER-SIMPLE-GUIDE.pptx"

MCP_JSON_WIN = """{
  "mcpServers": {
    "ifp-forecast": {
      "type": "stdio",
      "command": "py",
      "args": ["C:/Users/YourName/projects/ifp-mcp-server/server.py"],
      "env": {
        "IFP_API_URL": "http://inventory-forecasting-prod-dplus.ava.prod.hulu.com/v1/inventory/forecast",
        "IFP_SOURCE_HEADER": "RYM Frequency Cap Test",
        "IFP_TESTS_ROOT": "C:/Users/YourName/projects/ifp-frequency-cap-tests"
      }
    }
  }
}"""

MCP_JSON_MAC = """{
  "mcpServers": {
    "ifp-forecast": {
      "type": "stdio",
      "command": "python3",
      "args": ["/Users/YourName/projects/ifp-mcp-server/server.py"],
      "env": {
        "IFP_API_URL": "http://inventory-forecasting-prod-dplus.ava.prod.hulu.com/v1/inventory/forecast",
        "IFP_SOURCE_HEADER": "RYM Frequency Cap Test",
        "IFP_TESTS_ROOT": "/Users/YourName/projects/ifp-frequency-cap-tests"
      }
    }
  }
}"""


def _section(prs, title: str, subtitle: str) -> None:
    from build_deck import add_title_bar, blank_slide

    slide = blank_slide(prs)
    add_title_bar(slide, title, prs, subtitle)


def build(prs) -> None:
    from build_deck import (
        add_bullets,
        add_code_block,
        add_note,
        add_table,
        add_title_bar,
        blank_slide,
        ensure_pptx,
        title_slide,
    )

    ensure_pptx()

    title_slide(
        prs,
        "Cursor + IFP — New User Guide",
        "Simple setup for Windows and Mac",
        "Share with new hires · follow slides in order",
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Where Do I Start?", prs, "Pick the section you need")
    add_table(
        slide,
        ("I need to…", "Go to slides"),
        [
            ("Install Cursor", "Part A — Download Cursor"),
            ("Folders, Python, MCP", "Part B — Day 1 Setup"),
            ("MCP green, Run Mode", "Part C — Cursor Settings"),
            ("Agent, Run buttons", "Part D — Chat & Approvals"),
            ("First practice prompts", "Part E — Startup Prompts"),
            ("Run a forecast daily", "Part F — IFP Day to Day"),
            ("Copy portal targeting (F12)", "Part G — DevTools Lab"),
            ("Daily MSTR report", "Part H — MSTR"),
            ("Schedule MSTR emails", "Part I — Mission Control"),
            ("Something broke", "Part J — Fix It"),
        ],
        top=1.0,
        row_h=0.32,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Day 1 Checklist", prs, "Check off before live IFP forecasts")
    add_bullets(
        slide,
        [
            "□ Part A — Install Cursor",
            "□ Part B — Folders, Python, mcp.json, MCP green",
            "□ Part C — Settings (MCP + Run Mode)",
            "□ Part D — Read Agent vs Ask and Run buttons",
            "□ Part E — Do startup prompts (Ask then Agent)",
            "□ Part F — One test forecast (VPN on)",
            "□ Part H — Only if you own the daily report",
        ],
        size=15,
    )
    add_note(slide, "You do NOT need Node.js for this setup. Python 3.10+ is required.", top=4.2)

    # --- Part A ---
    _section(prs, "PART A — Download Cursor", "Install Cursor before anything else")

    slide = blank_slide(prs)
    add_title_bar(slide, "A1–A2 — Download Cursor (Windows)", prs, "cursor.com")
    add_bullets(
        slide,
        [
            "1. Open Chrome or Edge",
            "2. Address bar → type cursor.com → Enter",
            "3. Click Download → choose Windows",
            "4. Wait for Downloads\\Cursor Setup.exe",
        ],
        size=16,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "A1–A2 — Download Cursor (Mac)", prs, "cursor.com")
    add_bullets(
        slide,
        [
            "1. Open Safari or Chrome",
            "2. Address bar → type cursor.com → Enter",
            "3. Click Download → choose Mac",
            "4. Wait for Cursor.dmg in Downloads",
        ],
        size=16,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "A3–A5 — Install & First Open (Windows)", prs)
    add_bullets(
        slide,
        [
            "5. File Explorer → Downloads → double-click Cursor Setup.exe",
            "6. If asked “Allow changes?” → Yes",
            "7. Installer: Next → Install → Finish",
            "8. Start menu → search Cursor → open",
            "9. Sign in (work account if IT requires)",
            "10. Import VS Code settings? → Skip or Continue (either OK)",
            "11. Optional: pin Cursor to taskbar",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "A3–A5 — Install & First Open (Mac)", prs)
    add_bullets(
        slide,
        [
            "5. Downloads → double-click Cursor.dmg",
            "6. Drag Cursor into Applications",
            "7. Open Applications → Cursor (allow if macOS asks)",
            "8. Sign in (work account if IT requires)",
            "9. Import VS Code settings? → Skip or Continue (either OK)",
            "10. Optional: keep Cursor in Dock",
        ],
        size=14,
    )

    # --- Part B ---
    _section(prs, "PART B — Day 1 Setup", "One time only · ask a teammate for B4–B7 if stuck")

    slide = blank_slide(prs)
    add_title_bar(slide, "B1 — Check Access", prs)
    add_bullets(
        slide,
        [
            "1. Connect Hulu VPN — confirm it works",
            "2. Open IFP portal (SSO login):",
            "     https://ifp-portal-prod.aor.prod.hulu.com/home",
            "3. Ask trainer for git URLs or two zip files",
        ],
        size=15,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "B2–B3 — Folders You Need (Windows)", prs)
    add_code_block(
        slide,
        "C:\\Users\\YourName\\projects\\ifp-frequency-cap-tests\n"
        "C:\\Users\\YourName\\projects\\ifp-mcp-server",
        1.05,
        font_size=13,
    )
    add_bullets(
        slide,
        [
            "",
            "Create projects folder under your user folder if missing.",
            "Git (trainer gives URLs):",
        ],
        top=2.2,
        size=13,
    )
    add_code_block(
        slide,
        "cd C:\\Users\\YourName\\projects\ngit clone https://YOUR-TEAM-URL/ifp-frequency-cap-tests.git\ngit clone https://YOUR-TEAM-URL/ifp-mcp-server.git",
        3.0,
        font_size=11,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "B2–B3 — Folders You Need (Mac)", prs)
    add_code_block(
        slide,
        "/Users/YourName/projects/ifp-frequency-cap-tests\n/Users/YourName/projects/ifp-mcp-server",
        1.05,
        font_size=13,
    )
    add_bullets(
        slide,
        [
            "",
            "Create projects folder under your home folder if missing.",
            "Git (trainer gives URLs):",
        ],
        top=2.2,
        size=13,
    )
    add_code_block(
        slide,
        "cd /Users/YourName/projects\ngit clone https://YOUR-TEAM-URL/ifp-frequency-cap-tests.git\ngit clone https://YOUR-TEAM-URL/ifp-mcp-server.git",
        3.0,
        font_size=11,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "What Each Folder Is For", prs)
    add_table(
        slide,
        ("Folder", "You use it for…"),
        [
            ("ifp-frequency-cap-tests", "Open this in Cursor every day for IFP"),
            ("ifp-mcp-server", "server.py — Cursor starts it; you rarely open this folder"),
        ],
        top=1.2,
        row_h=0.55,
    )

    # Python — detailed Windows
    slide = blank_slide(prs)
    add_title_bar(slide, "B4 — Install Python (Windows)", prs, "Required for IFP MCP · Node.js NOT needed")
    add_bullets(
        slide,
        [
            "1. Browser → python.org/downloads",
            "2. Click Download Python 3.x (3.10 or newer)",
            "3. Run the .exe installer",
            "4. IMPORTANT: check Add python.exe to PATH at the bottom",
            "5. Click Install Now → Close when done",
            "6. You do NOT need the Python extension in Cursor for MCP",
        ],
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "B4 — Python Commands (Windows)", prs, "Run in PowerShell")
    add_code_block(
        slide,
        "py --version\n\n"
        "py -m pip install -r C:\\Users\\YourName\\projects\\ifp-mcp-server\\requirements.txt",
        1.05,
        font_size=12,
    )
    add_bullets(
        slide,
        [
            "",
            "Replace YourName with your Windows username.",
            "If py fails, ask trainer — PATH may not be set.",
        ],
        top=2.8,
        size=13,
    )

    # Python — detailed Mac
    slide = blank_slide(prs)
    add_title_bar(slide, "B4 — Install Python (Mac)", prs, "Required for IFP MCP · Node.js NOT needed")
    add_bullets(
        slide,
        [
            "Option A — python.org (recommended for new hires):",
            "  1. python.org/downloads → Download Python 3.x",
            "  2. Open the .pkg installer → Continue → Install",
            "  3. Enter Mac password if asked",
            "",
            "Option B — Homebrew (if you already use brew):",
            "  brew install python@3.12",
            "",
            "Use python3 (not the old system python). No Cursor Python extension required.",
        ],
        size=13,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "B4 — Python Commands (Mac)", prs, "Run in Terminal")
    add_code_block(
        slide,
        "python3 --version\n\n"
        "python3 -m pip install -r /Users/YourName/projects/ifp-mcp-server/requirements.txt",
        1.05,
        font_size=12,
    )
    add_bullets(
        slide,
        [
            "",
            "Replace YourName with your Mac username (whoami in Terminal).",
        ],
        top=2.6,
        size=13,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Windows vs Mac — Quick Reference", prs, "Same setup, different commands")
    add_table(
        slide,
        ("Task", "Windows", "Mac"),
        [
            ("Check Python", "py --version", "python3 --version"),
            ("Install MCP packages", "py -m pip install -r …", "python3 -m pip install -r …"),
            ("mcp.json command", '"command": "py"', '"command": "python3"'),
            ("Projects path", "C:/Users/You/projects/…", "/Users/You/projects/…"),
            ("MCP config file", "C:\\Users\\You\\.cursor\\mcp.json", "/Users/You/.cursor/mcp.json"),
            ("Terminal", "PowerShell", "Terminal (zsh)"),
        ],
        top=1.0,
        row_h=0.38,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "B5 — mcp.json (Windows)", prs, "Personal file — not in git")
    add_bullets(
        slide,
        [
            "1. C:\\Users\\YourName\\ → View → Hidden items",
            "2. Open or create folder .cursor",
            "3. Create file mcp.json — paste below",
            "4. Replace every YourName with your Windows username → Save",
        ],
        top=1.0,
        size=12,
    )
    add_code_block(slide, MCP_JSON_WIN, 2.35, font_size=8, max_height=4.0)

    slide = blank_slide(prs)
    add_title_bar(slide, "B5 — mcp.json (Mac)", prs, "Personal file — not in git")
    add_bullets(
        slide,
        [
            "1. Finder → Go → Home (or Cmd+Shift+H)",
            "2. Show hidden files: Cmd+Shift+. if needed",
            "3. Open or create folder .cursor",
            "4. Create file mcp.json — paste below",
            "5. Replace YourName with your Mac username → Save",
        ],
        top=1.0,
        size=12,
    )
    add_code_block(slide, MCP_JSON_MAC, 2.45, font_size=8, max_height=4.0)

    slide = blank_slide(prs)
    add_title_bar(slide, "B6–B7 — Reload & Open Project", prs)
    add_bullets(
        slide,
        [
            "1. Cursor: Ctrl+Shift+P (Mac: Cmd+Shift+P) → Reload Window → Enter",
            "2. File → Open Folder → ifp-frequency-cap-tests",
            "3. Sidebar should show config, docs, reference, scripts",
            "4. Part C → Settings → MCP → ifp-forecast must be GREEN",
        ],
        size=15,
    )

    # --- Part C ---
    _section(prs, "PART C — Cursor Settings", "MCP green + Run Mode")

    slide = blank_slide(prs)
    add_title_bar(slide, "C1–C2 — Open Settings & Check MCP", prs)
    add_bullets(
        slide,
        [
            "Open Settings: gear icon bottom-left OR Ctrl+, (Mac: Cmd+,)",
            "Search: MCP",
            "Find ifp-forecast → status must be GREEN",
            "If RED: click logs → ask trainer (wrong path, pip failed, bad JSON)",
        ],
        size=15,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "C3 — Run Mode (New Hires)", prs, "Pick a mode that asks before running things")
    add_table(
        slide,
        ("Run Mode", "What it means"),
        [
            ("Allowlist (empty)", "Asks before almost every command/MCP action — safest"),
            ("Auto-review", "Some things auto-run after review — fewer popups"),
            ("Run everything", "Do NOT use on day 1"),
        ],
        top=1.2,
        row_h=0.5,
    )
    add_bullets(slide, ["", "After any mcp.json edit: Reload Window"], top=3.5, size=14)

    # --- Part D ---
    _section(prs, "PART D — Chat & Approvals", "Ctrl+I · New Chat · Agent vs Ask")

    slide = blank_slide(prs)
    add_title_bar(slide, "Chat Basics", prs)
    add_table(
        slide,
        ("Mode", "When to use", "Runs IFP?"),
        [
            ("Ask", "Learning, explaining JSON", "No — safe to practice"),
            ("Agent", "Forecasts, scripts, edits", "Yes — with your approval"),
            ("Plan", "Big multi-step jobs", "Plan first, then Agent"),
        ],
        top=1.1,
        row_h=0.45,
    )
    add_bullets(
        slide,
        [
            "",
            "Ctrl+I (Mac: Cmd+I) → New Chat (+) for each new task",
            "IFP forecast → Agent · “What does this file mean?” → Ask",
        ],
        top=3.2,
        size=14,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "Run / Skip / Accept Buttons", prs)
    add_table(
        slide,
        ("Button", "What to do"),
        [
            ("Run / Approve", "OK once — VPN on for IFP; matches what you asked"),
            ("Skip / Reject", "Do not run — ask trainer"),
            ("Add to allowlist", "Ask trainer before using"),
            ("Accept (file edit)", "Keep the change"),
            ("Reject (file edit)", "Undo the change"),
        ],
        top=1.05,
        row_h=0.42,
    )

    # --- Part E ---
    _section(prs, "PART E — Startup Prompts", "Open ifp-frequency-cap-tests first")

    slide = blank_slide(prs)
    add_title_bar(slide, "E1 — Ask Mode (No VPN)", prs)
    add_code_block(
        slide,
        "Explain reference/cap-ratio-expectations.json in simple terms.\n"
        "What is one test case?\n\n"
        "What is the DMA code for Baton Rouge? Use reference/dma-codes.json.\n\n"
        "Explain config/base-request.json — dates and main targeting only.\n"
        "Do not change any files.",
        1.05,
        font_size=11,
    )

    slide = blank_slide(prs)
    add_title_bar(slide, "E2–E3 — Agent (VPN for Forecast)", prs)
    add_code_block(
        slide,
        "Create practice-notes.md with 4 bullets summarizing caps,\n"
        "DMA codes, and base-request.json.\n\n"
        "(VPN ON)\n"
        "Use ifp_server_info to show configuration.\n\n"
        "Run a forecast for July 8, 2026, Hulu US, Baton Rouge, 2 per hour.\n"
        "Show capacity and available.",
        1.05,
        font_size=11,
    )
    add_bullets(
        slide,
        ["Click Run on MCP approval cards · VPN must be on for forecasts"],
        top=4.2,
        size=14,
    )

    # --- Part F ---
    _section(prs, "PART F — IFP Day to Day", "Repeat every forecast")

    slide = blank_slide(prs)
    add_title_bar(slide, "Every Forecast", prs)
    add_table(
        slide,
        ("Step", "Do this"),
        [
            ("1", "Hulu VPN on"),
            ("2", "Open ifp-frequency-cap-tests in Cursor"),
            ("3", "Ctrl+I → New Chat → Agent"),
            ("4", "Plain English request with date, city, cap"),
            ("5", "Run on MCP approval if prompted"),
            ("6", "Read capacity and available in reply"),
        ],
        top=1.05,
        row_h=0.38,
    )

    # --- Part G ---
    _section(prs, "PART G — DevTools Lab", "When numbers ≠ portal")

    slide = blank_slide(prs)
    add_title_bar(slide, "Copy Portal Request (F12)", prs)
    add_bullets(
        slide,
        [
            "1. VPN on → IFP portal → build forecast like the deal",
            "2. F12 → Network → Preserve log → Fetch/XHR",
            "3. Run forecast in portal → click forecast row",
            "4. Copy Request JSON (NOT Response)",
            "5. Paste into config/base-request.json → Save → test Part F",
        ],
        size=14,
    )

    # --- Part H / I ---
    _section(prs, "PART H — MSTR Daily Report", "Different job — all-publishers-daily-report folder")

    slide = blank_slide(prs)
    add_title_bar(slide, "Daily Report Steps", prs)
    add_bullets(
        slide,
        [
            "1. Wait for MicroStrategy email in Outlook (.xlsx)",
            "2. Open all-publishers-daily-report in Cursor",
            "3. Ctrl+I → New Chat → Agent",
            "4. Type: Run dailypublisherscript.py",
            "5. Review Outlook draft → you click Send",
        ],
        size=15,
    )

    _section(prs, "PART I — Mission Control", "Schedule MSTR emails in browser — not Cursor")

    slide = blank_slide(prs)
    add_title_bar(slide, "Mission Control Setup", prs)
    add_bullets(
        slide,
        [
            "MicroStrategy Library → find All Publisher Daily Report",
            "Mission Control → Create subscription",
            "Recurrence: Daily · Format: Excel · Delivery: team email",
            "Save → test send → confirm email in Outlook",
        ],
        size=15,
    )

    # --- Part J / K ---
    _section(prs, "PART J — Fix It", "Common problems")

    slide = blank_slide(prs)
    add_title_bar(slide, "Troubleshooting", prs)
    add_table(
        slide,
        ("Problem", "Try this"),
        [
            ("Forecast fails", "VPN on?"),
            ("MCP red", "Check mcp.json paths + pip install + Reload Window"),
            ("Python not found (Win)", "Use py not python"),
            ("Python not found (Mac)", "Use python3 not python"),
            ("IFP ≠ portal", "Part G — update base-request.json targeting"),
            ("No MSTR email", "Part I — subscription failed or late"),
        ],
        top=1.05,
        row_h=0.42,
    )
    add_note(slide, "Never paste passwords or tokens in chat.", top=4.5, kind="tip")

    slide = blank_slide(prs)
    add_title_bar(slide, "Quick Reference", prs)
    add_table(
        slide,
        ("Item", "Windows", "Mac"),
        [
            ("Open chat", "Ctrl+I", "Cmd+I"),
            ("Settings", "Ctrl+,", "Cmd+,"),
            ("Reload", "Ctrl+Shift+P", "Cmd+Shift+P"),
            ("IFP folder", "...\\ifp-frequency-cap-tests", ".../ifp-frequency-cap-tests"),
            ("MCP config", "C:\\Users\\You\\.cursor\\mcp.json", "/Users/You/.cursor/mcp.json"),
        ],
        top=1.05,
        row_h=0.4,
    )


def main() -> Path:
    from build_deck import ensure_pptx

    ensure_pptx()
    from pptx import Presentation

    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    build(prs)
    out = OUTPUT.resolve()
    prs.save(out)
    return out


if __name__ == "__main__":
    path = main()
    print(f"Created: {path} ({path.stat().st_size // 1024} KB)")
